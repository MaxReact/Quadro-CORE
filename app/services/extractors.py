import json
import logging
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

logger = logging.getLogger(__name__)

MAX_CHARS = 30_000
_TRUNCATION_MARKER = "\n\n[...текст обрезан...]"
_XMIND_PARSE_ERROR = (
    "[Не удалось разобрать .xmind автоматически. "
    "Выгрузите карту как Markdown или текст и загрузите заново.]"
)


def _truncate(text: str) -> str:
    if len(text) > MAX_CHARS:
        return text[:MAX_CHARS] + _TRUNCATION_MARKER
    return text


# --- xlsx ---

def extract_xlsx(file_path: str) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    parts: list[str] = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        # Skip completely empty sheets
        if not any(any(cell is not None for cell in row) for row in rows):
            continue

        parts.append(f"## Лист: {sheet_name}")

        if not rows:
            continue

        headers = [str(c) if c is not None else "" for c in rows[0]]
        col_count = len(headers)

        # Markdown table header
        parts.append("| " + " | ".join(headers) + " |")
        parts.append("| " + " | ".join(["---"] * col_count) + " |")

        for row in rows[1:]:
            cells = [str(c) if c is not None else "" for c in row]
            # Pad or trim to column count
            while len(cells) < col_count:
                cells.append("")
            cells = cells[:col_count]
            parts.append("| " + " | ".join(cells) + " |")

        parts.append("")

    wb.close()
    return _truncate("\n".join(parts))


# --- docx ---

def extract_docx(file_path: str) -> str:
    from docx import Document
    from docx.oxml.ns import qn

    _HEADING_MAP = {
        "Heading 1": "#",
        "Heading 2": "##",
        "Heading 3": "###",
    }

    doc = Document(file_path)
    parts: list[str] = []

    def _table_to_md(table) -> str:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if not rows:
            return ""
        col_count = max(len(r) for r in rows)
        lines: list[str] = []
        header = rows[0] + [""] * (col_count - len(rows[0]))
        lines.append("| " + " | ".join(header) + " |")
        lines.append("| " + " | ".join(["---"] * col_count) + " |")
        for row in rows[1:]:
            cells = row + [""] * (col_count - len(row))
            lines.append("| " + " | ".join(cells) + " |")
        return "\n".join(lines)

    # Iterate body children in order to preserve paragraph/table interleaving
    body = doc.element.body
    for child in body:
        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag

        if tag == "p":
            from docx.text.paragraph import Paragraph
            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            style_name = para.style.name if para.style else ""
            prefix = _HEADING_MAP.get(style_name, "")
            parts.append(f"{prefix} {text}" if prefix else text)

        elif tag == "tbl":
            from docx.table import Table
            tbl = Table(child, doc)
            md_table = _table_to_md(tbl)
            if md_table:
                parts.append(md_table)
                parts.append("")

    return _truncate("\n".join(parts))


# --- xmind ---

def _parse_xmind_json(data: bytes) -> str:
    sheets = json.loads(data)
    parts: list[str] = []

    def _walk(topic: dict, depth: int) -> None:
        title = topic.get("title", "").strip()
        if depth == 0:
            parts.append(f"# {title}")
        else:
            parts.append("  " * (depth - 1) + f"- {title}")
        children = topic.get("children", {})
        attached = children.get("attached", []) if isinstance(children, dict) else []
        for child in attached:
            _walk(child, depth + 1)

    for sheet in sheets:
        root = sheet.get("rootTopic", {})
        _walk(root, 0)
        parts.append("")

    return "\n".join(parts)


def _parse_xmind_xml(data: bytes) -> str:
    root = ET.fromstring(data)
    ns = {"xmind": "urn:xmind:xmap:xmlns:content:2.0"}
    parts: list[str] = []

    # Try both namespaced and plain tag names
    def _find_topics(node):
        for tag in ("xmind:topic", "topic"):
            found = node.findall(f".//{tag}", ns) if ":" in tag else node.iter(tag)
            result = list(found)
            if result:
                return result
        return []

    def _walk(node, depth: int) -> None:
        title_node = (
            node.find("xmind:title", ns)
            or node.find("title")
        )
        title = (title_node.text or "").strip() if title_node is not None else ""
        if not title:
            return
        if depth == 0:
            parts.append(f"# {title}")
        else:
            parts.append("  " * (depth - 1) + f"- {title}")

        for children_tag in ("xmind:children", "children"):
            children_node = node.find(children_tag, ns) if ":" in children_tag else node.find(children_tag)
            if children_node is not None:
                for topic_tag in ("xmind:topic", "topic"):
                    for child in (children_node.findall(topic_tag, ns) if ":" in topic_tag else children_node.findall(topic_tag)):
                        _walk(child, depth + 1)
                break

    # Find sheet -> topic root
    for sheet_tag in ("xmind:sheet", "sheet"):
        sheets = root.findall(sheet_tag, ns) if ":" in sheet_tag else root.findall(sheet_tag)
        for sheet in sheets:
            for topic_tag in ("xmind:topic", "topic"):
                topics = sheet.findall(topic_tag, ns) if ":" in topic_tag else sheet.findall(topic_tag)
                for topic in topics:
                    _walk(topic, 0)
                    parts.append("")

    return "\n".join(parts)


def extract_xmind(file_path: str) -> str:
    try:
        with zipfile.ZipFile(file_path, "r") as zf:
            names = zf.namelist()
            if "content.json" in names:
                return _truncate(_parse_xmind_json(zf.read("content.json")))
            elif "content.xml" in names:
                return _truncate(_parse_xmind_xml(zf.read("content.xml")))
            else:
                logger.warning("xmind: neither content.json nor content.xml found in %s", file_path)
                return _XMIND_PARSE_ERROR
    except Exception as exc:
        logger.warning("xmind parse error for %s: %s", file_path, exc)
        return _XMIND_PARSE_ERROR


# --- txt / md ---

def extract_text(file_path: str) -> str:
    text = Path(file_path).read_text(encoding="utf-8", errors="replace")
    return _truncate(text)


# --- Dispatch ---

_EXTRACTORS = {
    "xlsx": extract_xlsx,
    "docx": extract_docx,
    "xmind": extract_xmind,
    "txt": extract_text,
    "md": extract_text,
}


def extract(file_path: str, ext: str) -> str:
    fn = _EXTRACTORS.get(ext.lower())
    if fn is None:
        return f"[Формат {ext!r} не поддерживается]"
    return fn(file_path)
